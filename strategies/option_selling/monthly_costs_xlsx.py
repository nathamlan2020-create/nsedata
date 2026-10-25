import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter as L
SP = "/tmp/claude-0/-home-claude/65f2a500-4991-5704-937b-9267d8da030c/scratchpad"
D = pd.read_pickle(f"{SP}/monthly_costs.pkl"); G1, G2 = D["G1"], D["G2"]
wb = Workbook(); F = "Arial"; RS = '#,##0;[Red]-#,##0;-'
H = Font(name=F, bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="1F4E78"); B = Font(name=F, bold=True); N = Font(name=F)
red = PatternFill("solid", fgColor="F8D7DA"); green = PatternFill("solid", fgColor="D4EDDA")
def head(ws, r, vals):
    for j, v in enumerate(vals, 1):
        c = ws.cell(r, j, v); c.font = H; c.fill = HF; c.alignment = Alignment(wrap_text=True, horizontal="center")
def sheet(name, df, cols, gross_cols, cost_cols, notes):
    ws = wb.create_sheet(name); head(ws, 1, ["Month"] + cols + ["Gross total ₹", "Costs total ₹", "Net total ₹ (after costs)", "Cumulative net ₹"])
    for i, (m, r) in enumerate(df.iterrows(), start=2):
        ws.cell(i, 1, m.strftime("%b %Y")).font = B
        for j, c in enumerate(cols, start=2): ws.cell(i, j, float(round(r[c]))).number_format = RS
        g = len(cols) + 2
        ws.cell(i, g, "=" + "+".join(f"{L(cols.index(c)+2)}{i}" for c in gross_cols)).number_format = RS
        ws.cell(i, g + 1, "=" + "+".join(f"{L(cols.index(c)+2)}{i}" for c in cost_cols)).number_format = RS
        ws.cell(i, g + 2, f"={L(g)}{i}-{L(g+1)}{i}").number_format = RS; ws.cell(i, g + 2).font = B
        ws.cell(i, g + 3, f"={L(g+2)}{i}" if i == 2 else f"={L(g+3)}{i-1}+{L(g+2)}{i}").number_format = RS
        for j in range(2, g + 4):
            if ws.cell(i, j).font != B: ws.cell(i, j).font = N
    last = i; nc = f"{L(len(cols)+4)}"
    ws.conditional_formatting.add(f"{nc}2:{nc}{last}", CellIsRule(operator="lessThan", formula=["0"], fill=red))
    ws.conditional_formatting.add(f"{nc}2:{nc}{last}", CellIsRule(operator="greaterThan", formula=["0"], fill=green))
    t = last + 1; ws.cell(t, 1, "Total").font = B
    for j in range(2, len(cols) + 5): ws.cell(t, j, f"=SUM({L(j)}2:{L(j)}{last})").number_format = RS; ws.cell(t, j).font = B
    for k, n in enumerate(notes): ws.cell(t + 2 + k, 1, n).font = Font(name=F, italic=True, size=9)
    ws.freeze_panes = "B2"; ws.column_dimensions["A"].width = 11
    for j in range(2, len(cols) + 6): ws.column_dimensions[L(j)].width = 14
    ws.row_dimensions[1].height = 45
    return ws, last
common = ["All P&L is after costs where marked 'net'. Costs = STT 0.1% on sell premium + ~0.04% exchange/SEBI/GST/stamp per side + ₹40 brokerage per leg round trip + slippage (Nifty 0.5 pt, Sensex 1.5 pt per side; +1% on stop fills).",
          "Option prices are rebuilt minute by minute (Nifty: calibrated to the real 9:15 bhavcopy straddle; Sensex: India-VIX based, Nifty path stretched to Sensex's real daily range). Bank Nifty uses real NSE daily prices (9:15 open entry, close exit). Treat ₹ figures as estimates.",
          "Income tax not included: F&O profit is business income taxed at your slab rate. Lots: Nifty 65, Sensex 20, Bank Nifty 30 (today's sizes) in every year."]
sheet("Nifty 3 lots monthly", G1, ["trades", "gross", "costs"], ["gross"], ["costs"],
      ["Nifty Tuesday/weekly expiry: 3 lots at 9:22 / 9:35 / 10:05, 40% SL per leg, other leg SL to cost, one re-entry before 1 PM, exit 15:15. 'trades' = expiry days in the month."] + common)
cols = ["gross_Nifty", "costs_Nifty", "gross_Sensex", "costs_Sensex", "gross_BankNifty", "costs_BankNifty"]
G2b = G2.copy()
ws, last = sheet("Weekly plan monthly", G2b, cols, ["gross_Nifty", "gross_Sensex", "gross_BankNifty"], ["costs_Nifty", "costs_Sensex", "costs_BankNifty"],
      ["Weekly plan since Sensex weeklies began (May 2023): Nifty 3 lots (as above) + Sensex 1 lot at 10:05 with 30% SL per leg + Bank Nifty 1 lot on the monthly expiry (open entry, 30% SL per leg)."] + common)
hdr = ["Month", "Nifty gross ₹", "Nifty costs ₹", "Sensex gross ₹", "Sensex costs ₹", "Bank Nifty gross ₹", "Bank Nifty costs ₹"]
for j, v in enumerate(hdr, 1): ws.cell(1, j, v)
w0 = wb.worksheets[1]
for j, v in enumerate(["Month", "Expiry days", "Gross ₹", "Costs ₹"], 1): w0.cell(1, j, v)
# yearly sheet
wy = wb.create_sheet("Yearly"); head(wy, 1, ["Year", "Nifty gross ₹", "Nifty costs ₹", "Nifty net ₹", "Plan gross ₹", "Plan costs ₹", "Plan net ₹ (Nifty+Sensex+BankNifty)"])
Y1 = G1.groupby(G1.index.year)[["gross", "costs", "net"]].sum(); Y2 = G2.groupby(G2.index.year)[["gross_total", "costs_total", "net_total"]].sum()
for i, y in enumerate(Y1.index, start=2):
    wy.cell(i, 1, str(y)).font = B
    wy.cell(i, 2, float(round(Y1.loc[y, "gross"]))); wy.cell(i, 3, float(round(Y1.loc[y, "costs"]))); wy.cell(i, 4, f"=B{i}-C{i}")
    if y in Y2.index: wy.cell(i, 5, float(round(Y2.loc[y, "gross_total"]))); wy.cell(i, 6, float(round(Y2.loc[y, "costs_total"]))); wy.cell(i, 7, f"=E{i}-F{i}")
    for j in range(2, 8): wy.cell(i, j).number_format = RS; wy.cell(i, j).font = N
wy.cell(i + 2, 1, "2019 = weekly Nifty expiries from Feb; 2023 plan = May-Dec; 2026 = to 23 Oct (Nifty) / 19 Oct (Sensex).").font = Font(name=F, italic=True, size=9)
for j in range(1, 8): wy.column_dimensions[L(j)].width = 16
wb.remove(wb["Sheet"]); wb.save("options_plan_monthly_after_costs.xlsx")
